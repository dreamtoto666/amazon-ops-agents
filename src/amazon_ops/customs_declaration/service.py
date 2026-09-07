from __future__ import annotations

import io
import json
import re
import zipfile
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from importlib.resources import files
from pathlib import Path
from typing import Any, Iterable, Mapping
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

from .models import (
    AggregatedItem,
    CategoryPreview,
    CategoryResult,
    CustomsIssue,
    CustomsPreview,
    PreviewItem,
    ProductRule,
    ShippingSpeed,
    TradeMode,
)


MAX_FILE_SIZE = 20 * 1024 * 1024
MAX_EXPANDED_SIZE = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 2_000
MAX_SHEET_ROWS = 50_000
TEMPLATE_CAPACITIES: Mapping[TradeMode, int] = {
    TradeMode.GENERAL: 23,
    TradeMode.CROSS_BORDER: 16,
}
MONEY_QUANTUM = Decimal("0.01")
WEIGHT_QUANTUM = Decimal("0.01")

_SHIPPING_CHANNELS: Mapping[ShippingSpeed, tuple[str, ...]] = {
    ShippingSpeed.STANDARD: (
        "美西普船快递派统配",
        "美西普船卡派统配（包）",
        "美东普船限时达纽约卡派（包）",
    ),
    ShippingSpeed.EXPRESS: (
        "美东快船海铁纽约卡派（包）",
        "美西快船卡派（包）",
        "美西快船快递派",
    ),
    ShippingSpeed.MATSON: (
        "美森极致达卡派（包）",
        "美森极致达快递派",
    ),
    ShippingSpeed.STANDARD_ECONOMY: (
        "美东普船限时达卡派特惠（包）",
        "美西普船卡派特惠（包）",
    ),
}

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", _MAIN_NS)
ET.register_namespace("r", _REL_NS)

_TYPE_ALIASES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("车库门密封条套装", ("车库门密封条套装", "带导轨门底密封条", "导轨")),
    ("车库门侧密封条", ("车库门侧密封条", "l型密封条", "l型")),
    ("U型边缘装饰条", ("u型边缘装饰条", "u型条", "边缘装饰条")),
    ("五孔形车用密封条", ("五孔形车用密封条", "五孔形", "五孔")),
    ("六孔形车用密封条", ("六孔形车用密封条", "六孔形", "六孔")),
    ("发泡形车用密封条", ("发泡形车用密封条", "发泡形", "发泡")),
    ("眼睛形车用密封条", ("眼睛形车用密封条", "眼睛形", "眼睛")),
    ("葫芦形门窗密封条", ("葫芦形门窗密封条", "葫芦形", "葫芦")),
    ("D形门窗密封条", ("d形门窗密封条", "d型", "工字型")),
    ("扇形墙角装饰条", ("扇形墙角装饰条", "扇形双面装饰条", "扇形")),
    ("墙角保护条", ("墙角保护条", "墙角保护器", "平面墙角保护器")),
    ("角落装饰条", ("角落装饰条", "防霉条")),
    ("柔性踢脚线", ("柔性踢脚线", "踢脚线")),
    ("镂空挡水条", ("镂空挡水条", "镂空挡水")),
    ("T型挡水条", ("t型挡水条", "t型挡水")),
    ("挡水条", ("小挡水条", "挡水条")),
    ("车库门密封条", ("车库门密封条", "u+o泡胶条", "u+o")),
)

# 已确认的供应商品名/长度别名，统一回落到已有报关商品规则。
_PRODUCT_NUMBER_ALIASES: Mapping[tuple[str, tuple[str, ...]], tuple[str, ...]] = {
    ("发泡形车用密封条", ("10", "10", "3.05")): ("10", "10", "3"),
    ("发泡形车用密封条", ("10", "10", "6.1")): ("10", "10", "6"),
    ("发泡形车用密封条", ("15", "15", "3.05")): ("15", "15", "3"),
    ("发泡形车用密封条", ("15", "15", "6.1")): ("15", "15", "6"),
    ("发泡形车用密封条", ("19", "19", "3.05")): ("19", "19", "3"),
    ("发泡形车用密封条", ("19", "19", "6.1")): ("19", "19", "6"),
    ("五孔形车用密封条", ("15", "12", "5.03")): ("5.03",),
    ("五孔形车用密封条", ("15", "12", "10.36")): ("10.36",),
    ("五孔形车用密封条", ("15", "12", "16.4")): ("16.4",),
    ("五孔形车用密封条", ("15", "12", "21.5")): ("21.5",),
    ("五孔形车用密封条", ("15", "12", "26.5")): ("26.5",),
}


@dataclass(slots=True)
class _Analysis:
    preview: CustomsPreview
    results: list[CategoryResult]


@dataclass(slots=True)
class _ShipmentAccumulator:
    box_count: int
    gross_weight: Decimal


class CustomsGenerationError(ValueError):
    def __init__(self, preview: CustomsPreview):
        super().__init__("报关数据校验未通过。")
        self.preview = preview


class CustomsDeclarationService:
    """Deterministic XLSX-to-customs-declaration processing service."""

    def __init__(
        self,
        *,
        catalog_path: Path | None = None,
        template_path: Path | None = None,
        template_paths: Mapping[TradeMode, Path] | None = None,
    ) -> None:
        resource_root = files("amazon_ops.customs_declaration.resources")
        self.catalog_path = catalog_path or Path(str(resource_root.joinpath("product_catalog.v1.json")))
        if template_paths is not None:
            self.template_paths = dict(template_paths)
        elif template_path is not None:
            # 保留测试及旧调用方传入单模板的兼容性。
            self.template_paths = {mode: template_path for mode in TradeMode}
        else:
            self.template_paths = {
                TradeMode.GENERAL: Path(
                    str(resource_root.joinpath("customs_declaration_0110.v1.xlsx"))
                ),
                TradeMode.CROSS_BORDER: Path(
                    str(resource_root.joinpath("customs_declaration_9810.v1.xlsx"))
                ),
            }
        missing_templates = set(TradeMode) - set(self.template_paths)
        if missing_templates:
            names = "、".join(mode.value for mode in sorted(missing_templates, key=str))
            raise ValueError(f"缺少贸易方式 {names} 的报关单模板。")
        # `template_path` 保留为 0110 模板，兼容既有集成与测试代码。
        self.template_path = self.template_paths[TradeMode.GENERAL]
        self.rules = self._load_catalog(self.catalog_path)
        self.rules_by_type: dict[str, list[ProductRule]] = {}
        self.rules_by_sku: dict[str, set[str]] = {}
        for rule in self.rules:
            self.rules_by_type.setdefault(rule.declaration_name, []).append(rule)
            for sku in rule.skus:
                self.rules_by_sku.setdefault(sku.upper(), set()).add(rule.rule_id)
        self.contract_numbers = {
            mode: self._read_template_contract_number(path)
            for mode, path in self.template_paths.items()
        }
        self.contract_number = self.contract_numbers[TradeMode.GENERAL]

    def preview(
        self,
        *,
        shipment_file_name: str,
        shipment_content: bytes,
        fba_file_name: str,
        fba_content: bytes,
        declaration_date: date | None = None,
    ) -> CustomsPreview:
        return self._analyze(
            shipment_file_name=shipment_file_name,
            shipment_content=shipment_content,
            fba_file_name=fba_file_name,
            fba_content=fba_content,
            declaration_date=declaration_date or self._shanghai_today(),
        ).preview

    def generate(
        self,
        *,
        shipment_file_name: str,
        shipment_content: bytes,
        fba_file_name: str,
        fba_content: bytes,
        declaration_date: date | None = None,
    ) -> tuple[str, bytes]:
        selected_date = declaration_date or self._shanghai_today()
        analysis = self._analyze(
            shipment_file_name=shipment_file_name,
            shipment_content=shipment_content,
            fba_file_name=fba_file_name,
            fba_content=fba_content,
            declaration_date=selected_date,
        )
        if not analysis.preview.can_generate:
            raise CustomsGenerationError(analysis.preview)

        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for result in analysis.results:
                if not result.items:
                    continue
                workbook = self._render_workbook(result, selected_date)
                integer_amount = result.amount.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
                file_name = (
                    f"{self.contract_numbers[result.trade_mode]}-美国海运-{result.shipping_speed.value}-"
                    f"{result.trade_mode.value}-{result.box_count}箱-报关单-"
                    f"{integer_amount}美元.xlsx"
                )
                archive.writestr(file_name, workbook)
        return f"{selected_date:%Y%m%d}-报关单.zip", output.getvalue()

    def _analyze(
        self,
        *,
        shipment_file_name: str,
        shipment_content: bytes,
        fba_file_name: str,
        fba_content: bytes,
        declaration_date: date,
    ) -> _Analysis:
        errors: list[CustomsIssue] = []
        warnings: list[CustomsIssue] = []
        assignments = self._parse_shipment_assignments(
            shipment_file_name, shipment_content, errors
        )
        raw_categories = self._parse_fba(
            fba_file_name,
            fba_content,
            assignments,
            errors,
            warnings,
        )

        results: list[CategoryResult] = []
        category_previews: list[CategoryPreview] = []
        for speed, trade in self._category_keys():
            order_ids = [
                order_id
                for order_id, category in assignments.items()
                if category == (speed, trade)
            ]
            result = self._aggregate_category(
                speed,
                trade,
                order_ids,
                raw_categories.get((speed, trade)),
                errors,
            )
            results.append(result)
            category_previews.append(self._to_preview(result))

        can_generate = not errors and any(result.items for result in results)
        return _Analysis(
            preview=CustomsPreview(
                declaration_date=declaration_date,
                can_generate=can_generate,
                categories=category_previews,
                warnings=warnings,
                errors=errors,
            ),
            results=results,
        )

    def _parse_shipment_assignments(
        self,
        file_name: str,
        content: bytes,
        errors: list[CustomsIssue],
    ) -> OrderedDict[str, tuple[ShippingSpeed, TradeMode]]:
        assignments: OrderedDict[str, tuple[ShippingSpeed, TradeMode]] = OrderedDict()
        workbook = self._open_input_workbook(file_name, content, errors, "一八发货模板")
        if workbook is None:
            return assignments
        if "美国专线箱单" not in workbook.sheetnames:
            errors.append(
                CustomsIssue(
                    code="MISSING_SHEET",
                    message="一八发货模板缺少“美国专线箱单”工作表。",
                    file=file_name,
                    sheet="美国专线箱单",
                )
            )
            return assignments
        sheet = workbook["美国专线箱单"]
        if sheet.max_row > MAX_SHEET_ROWS:
            errors.append(self._row_limit_issue(file_name, sheet.title))
            return assignments
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header_map = {self._header_key(value): index for index, value in enumerate(header)}
        route_index = self._find_header_index(header_map, "走货渠道")
        trade_index = self._find_header_index(header_map, "贸易方式")
        if len(header) < 4 or route_index is None or trade_index is None:
            errors.append(
                CustomsIssue(
                    code="MISSING_HEADER",
                    message="美国专线箱单缺少客户原单号、走货渠道或贸易方式表头。",
                    file=file_name,
                    sheet=sheet.title,
                )
            )
            return assignments

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            order_id = self._text(row[3] if len(row) > 3 else None)
            if not order_id:
                continue
            route = self._text(row[route_index] if len(row) > route_index else None)
            speed = self._shipping_speed(route)
            if speed is None:
                errors.append(
                    CustomsIssue(
                        code="INVALID_SHIPPING_CHANNEL",
                        message=(
                            f"客户原单号 {order_id} 的走货渠道未匹配普船统配、快船、"
                            "美森极致达或普船特惠。"
                        ),
                        file=file_name,
                        sheet=sheet.title,
                        row=row_number,
                        field="走货渠道",
                    )
                )
                continue
            trade = self._trade_mode(row[trade_index] if len(row) > trade_index else None)
            if trade is None:
                errors.append(
                    CustomsIssue(
                        code="INVALID_TRADE_MODE",
                        message=f"客户原单号 {order_id} 的贸易方式必须是 0110 或 9810。",
                        file=file_name,
                        sheet=sheet.title,
                        row=row_number,
                        field="贸易方式",
                    )
                )
                continue
            category = (speed, trade)
            previous = assignments.get(order_id)
            if previous is not None and previous != category:
                errors.append(
                    CustomsIssue(
                        code="ORDER_CATEGORY_CONFLICT",
                        message=f"客户原单号 {order_id} 同时出现在不同报关分类。",
                        file=file_name,
                        sheet=sheet.title,
                        row=row_number,
                        field="客户原单号",
                    )
                )
                continue
            assignments.setdefault(order_id, category)
        return assignments

    def _parse_fba(
        self,
        file_name: str,
        content: bytes,
        assignments: OrderedDict[str, tuple[ShippingSpeed, TradeMode]],
        errors: list[CustomsIssue],
        warnings: list[CustomsIssue],
    ) -> dict[tuple[ShippingSpeed, TradeMode], dict[str, Any]]:
        categories: dict[tuple[ShippingSpeed, TradeMode], dict[str, Any]] = {}
        workbook = self._open_input_workbook(file_name, content, errors, "FBA货件表")
        if workbook is None:
            return categories
        if "装箱明细" not in workbook.sheetnames:
            errors.append(
                CustomsIssue(
                    code="MISSING_SHEET",
                    message="FBA货件表缺少“装箱明细”工作表。",
                    file=file_name,
                    sheet="装箱明细",
                )
            )
            return categories
        detail_sheet_name = next(
            (name for name in ("货件详情", "货件详细") if name in workbook.sheetnames),
            None,
        )
        if detail_sheet_name is None:
            errors.append(
                CustomsIssue(
                    code="MISSING_SHEET",
                    message="FBA货件表缺少“货件详情”工作表。",
                    file=file_name,
                    sheet="货件详情",
                )
            )
            return categories

        packing_sheet = workbook["装箱明细"]
        detail_sheet = workbook[detail_sheet_name]
        if packing_sheet.max_row > MAX_SHEET_ROWS:
            errors.append(self._row_limit_issue(file_name, packing_sheet.title))
            return categories
        if detail_sheet.max_row > MAX_SHEET_ROWS:
            errors.append(self._row_limit_issue(file_name, detail_sheet.title))
            return categories

        packing_header = next(
            packing_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
        )
        packing_indices = self._required_header_indices(
            packing_header, ("货件单号", "总箱数", "总重量")
        )
        if packing_indices is None:
            errors.append(
                CustomsIssue(
                    code="MISSING_HEADER",
                    message="装箱明细缺少货件单号、总箱数或总重量表头。",
                    file=file_name,
                    sheet=packing_sheet.title,
                )
            )
            return categories

        detail_header = next(
            detail_sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
        )
        detail_indices = self._required_header_indices(
            detail_header, ("货件单号", "品名", "SKU", "申报量")
        )
        if detail_indices is None:
            errors.append(
                CustomsIssue(
                    code="MISSING_HEADER",
                    message="货件详情缺少货件单号、品名、SKU或申报量表头。",
                    file=file_name,
                    sheet=detail_sheet.title,
                )
            )
            return categories

        current_order: str | None = None
        seen_shipments: set[str] = set()
        packed_shipments: set[str] = set()

        for row_number, row in enumerate(packing_sheet.iter_rows(min_row=2, values_only=True), 2):
            new_order = self._value_at(row, packing_indices["货件单号"])
            new_order = self._text(new_order)
            if new_order:
                if new_order in seen_shipments:
                    errors.append(
                        CustomsIssue(
                            code="DUPLICATE_SHIPMENT_GROUP",
                            message=f"货件单号 {new_order} 在装箱明细中出现了多个分组。",
                            file=file_name,
                            sheet=packing_sheet.title,
                            row=row_number,
                            field="货件单号",
                        )
                    )
                seen_shipments.add(new_order)
                current_order = new_order
                if current_order in assignments:
                    category = assignments[current_order]
                    box_count = self._positive_integer(
                        self._value_at(row, packing_indices["总箱数"]),
                        errors,
                        file_name,
                        packing_sheet.title,
                        row_number,
                        "总箱数",
                    )
                    gross_weight = self._positive_decimal(
                        self._value_at(row, packing_indices["总重量"]),
                        errors,
                        file_name,
                        packing_sheet.title,
                        row_number,
                        "总重量",
                    )
                    bucket = categories.setdefault(
                        category,
                        {"shipments": {}, "quantities": OrderedDict()},
                    )
                    if box_count is not None and gross_weight is not None:
                        bucket["shipments"][current_order] = _ShipmentAccumulator(
                            box_count=box_count,
                            gross_weight=gross_weight,
                        )
            elif current_order is None:
                continue
            if current_order in assignments:
                packed_shipments.add(current_order)

        current_order = None
        detailed_shipments: set[str] = set()
        for row_number, row in enumerate(detail_sheet.iter_rows(min_row=2, values_only=True), 2):
            new_order = self._text(self._value_at(row, detail_indices["货件单号"]))
            if new_order:
                current_order = new_order
            elif current_order is None:
                continue
            if current_order not in assignments:
                continue

            product_name = self._text(self._value_at(row, detail_indices["品名"]))
            sku = self._text(self._value_at(row, detail_indices["SKU"])).upper()
            quantity = self._positive_decimal(
                self._value_at(row, detail_indices["申报量"]),
                errors,
                file_name,
                detail_sheet.title,
                row_number,
                "申报量",
            )
            if not product_name or quantity is None:
                if not product_name:
                    errors.append(
                        CustomsIssue(
                            code="MISSING_PRODUCT_NAME",
                            message="货件详情中的品名不能为空。",
                            file=file_name,
                            sheet=detail_sheet.title,
                            row=row_number,
                            field="品名",
                        )
                    )
                continue
            # 货件已有有效商品行；后续的未知品名不应再被误报为“找不到商品明细”。
            detailed_shipments.add(current_order)
            rule = self._match_product(product_name)
            if rule is None:
                errors.append(
                    CustomsIssue(
                        code="UNKNOWN_PRODUCT",
                        message=f"品名“{product_name}”未匹配到内置商品目录。",
                        file=file_name,
                        sheet=detail_sheet.title,
                        row=row_number,
                        field="品名",
                    )
                )
                continue
            if rule.unit_weight <= 0 or rule.unit_price < 0:
                errors.append(
                    CustomsIssue(
                        code="INCOMPLETE_PRODUCT_RULE",
                        message=f"商品规则“{rule.standard_name}”缺少有效重量或价格。",
                        file=file_name,
                        sheet=detail_sheet.title,
                        row=row_number,
                        field="品名",
                    )
                )
                continue
            sku_rules = self.rules_by_sku.get(sku)
            if sku and sku_rules and rule.rule_id not in sku_rules:
                warnings.append(
                    CustomsIssue(
                        code="SKU_NAME_CONFLICT",
                        message=(
                            f"SKU {sku} 与品名“{product_name}”对应规则不一致，"
                            "已按品名统计。"
                        ),
                        file=file_name,
                        sheet=detail_sheet.title,
                        row=row_number,
                        field="SKU",
                    )
                )
            category = assignments[current_order]
            bucket = categories.setdefault(
                category,
                {"shipments": {}, "quantities": OrderedDict()},
            )
            quantities: OrderedDict[str, Decimal] = bucket["quantities"]
            quantities[rule.rule_id] = quantities.get(rule.rule_id, Decimal("0")) + quantity

        for order_id in assignments:
            if order_id not in packed_shipments:
                errors.append(
                    CustomsIssue(
                        code="MISSING_SHIPMENT",
                        message=f"FBA装箱明细中找不到客户原单号 {order_id}。",
                        file=file_name,
                        sheet=packing_sheet.title,
                        field="货件单号",
                    )
                )
            if order_id not in detailed_shipments:
                errors.append(
                    CustomsIssue(
                        code="MISSING_SHIPMENT_DETAIL",
                        message=f"FBA货件详情中找不到客户原单号 {order_id} 的商品明细。",
                        file=file_name,
                        sheet=detail_sheet.title,
                        field="货件单号",
                    )
                )
        return categories

    def _aggregate_category(
        self,
        speed: ShippingSpeed,
        trade: TradeMode,
        order_ids: list[str],
        raw: dict[str, Any] | None,
        errors: list[CustomsIssue],
    ) -> CategoryResult:
        if raw is None:
            return CategoryResult(
                shipping_speed=speed,
                trade_mode=trade,
                order_ids=order_ids,
                shipment_count=0,
                box_count=0,
                gross_weight=Decimal("0"),
                net_weight=Decimal("0"),
                amount=Decimal("0"),
                items=[],
            )
        shipments: dict[str, _ShipmentAccumulator] = raw["shipments"]
        quantities: OrderedDict[str, Decimal] = raw["quantities"]
        items: list[AggregatedItem] = []
        for rule in self.rules:
            quantity = quantities.get(rule.rule_id)
            if quantity is None:
                continue
            amount = (quantity * rule.unit_price).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP)
            items.append(
                AggregatedItem(
                    rule=rule,
                    quantity=quantity,
                    net_weight=Decimal("0"),
                    gross_weight=Decimal("0"),
                    amount=amount,
                )
            )
        capacity = TEMPLATE_CAPACITIES[trade]
        if len(items) > capacity:
            errors.append(
                CustomsIssue(
                    code="TEMPLATE_CAPACITY_EXCEEDED",
                    message=(
                        f"{speed.value} × {trade.value} 有 {len(items)} 个标准商品，"
                        f"超过 {trade.value} 模板最多 {capacity} 个的容量。"
                    ),
                )
            )
        box_count = sum(item.box_count for item in shipments.values())
        gross_weight = sum(
            (item.gross_weight for item in shipments.values()), Decimal("0")
        ).quantize(WEIGHT_QUANTUM, rounding=ROUND_HALF_UP)
        theoretical_weight = sum(
            (item.quantity * item.rule.unit_weight for item in items), Decimal("0")
        )
        net_weight = (gross_weight * Decimal("0.92")).quantize(
            WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
        )
        amount = sum((item.amount for item in items), Decimal("0")).quantize(
            MONEY_QUANTUM, rounding=ROUND_HALF_UP
        )
        if items and theoretical_weight <= 0:
            errors.append(
                CustomsIssue(
                    code="ZERO_NET_WEIGHT",
                    message=f"{speed.value} × {trade.value} 缺少可用于分配重量的商品单件重量。",
                )
            )
        elif items:
            allocated_gross = Decimal("0")
            allocated_net = Decimal("0")
            for item in items[:-1]:
                weight_ratio = item.quantity * item.rule.unit_weight / theoretical_weight
                item.gross_weight = (gross_weight * weight_ratio).quantize(
                    WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
                )
                item.net_weight = (net_weight * weight_ratio).quantize(
                    WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
                )
                allocated_gross += item.gross_weight
                allocated_net += item.net_weight
            items[-1].gross_weight = (gross_weight - allocated_gross).quantize(
                WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
            )
            items[-1].net_weight = (net_weight - allocated_net).quantize(
                WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
            )
        return CategoryResult(
            shipping_speed=speed,
            trade_mode=trade,
            order_ids=order_ids,
            shipment_count=len(shipments),
            box_count=box_count,
            gross_weight=gross_weight,
            net_weight=net_weight,
            amount=amount,
            items=items,
        )

    def _to_preview(self, result: CategoryResult) -> CategoryPreview:
        items = [
            PreviewItem(
                rule_id=item.rule.rule_id,
                declaration_name=item.rule.declaration_name,
                model=item.rule.model,
                material=item.rule.material,
                hs_code=item.rule.hs_code,
                quantity=item.quantity,
                unit_weight=item.rule.unit_weight,
                net_weight=item.net_weight,
                gross_weight=item.gross_weight,
                unit_price=item.rule.unit_price,
                amount=item.amount,
            )
            for item in result.items
        ]
        return CategoryPreview(
            shipping_speed=result.shipping_speed,
            trade_mode=result.trade_mode,
            order_count=len(result.order_ids),
            shipment_count=result.shipment_count,
            item_count=len(items),
            box_count=result.box_count,
            net_weight=result.net_weight,
            gross_weight=result.gross_weight,
            amount=result.amount,
            generates_file=bool(items),
            items=items,
        )

    def _match_product(self, product_name: str) -> ProductRule | None:
        cleaned = self._normalize_name(product_name)
        product_type = self._detect_product_type(cleaned)
        if product_type is None:
            return None
        numbers = self._number_key(cleaned)
        numbers = _PRODUCT_NUMBER_ALIASES.get((product_type, numbers), numbers)
        candidates = self.rules_by_type.get(product_type, [])
        exact = [rule for rule in candidates if rule.number_key == numbers]
        return exact[0] if len(exact) == 1 else None

    def _render_workbook(
        self,
        result: CategoryResult,
        declaration_date: date,
        *,
        destination_country: str = "美国",
        trade_country: str | None = None,
        arrival_country: str | None = None,
        destination_port: str | None = None,
        transportation_mode: str | None = None,
    ) -> bytes:
        template = self.template_paths[result.trade_mode].read_bytes()
        source = zipfile.ZipFile(io.BytesIO(template), "r")
        sheet_path = self._report_sheet_path(source)
        sheet_xml = source.read(sheet_path)
        root = ET.fromstring(sheet_xml)

        self._normalize_item_styles(root, TEMPLATE_CAPACITIES[result.trade_mode])
        self._set_cell(root, "K4", to_excel(datetime.combine(declaration_date, datetime.min.time())))
        if trade_country is not None:
            self._set_cell(root, "E10", trade_country)
        if arrival_country is not None:
            self._set_cell(root, "G10", arrival_country)
        if destination_port is not None:
            self._set_cell(root, "K10", destination_port)
        if transportation_mode is not None:
            self._set_cell(root, "E6", transportation_mode)
        for index, item in enumerate(result.items):
            base = 20 + index * 3
            self._set_cell(root, f"B{base}", item.rule.hs_code)
            self._set_cell(root, f"D{base}", item.rule.declaration_name)
            self._set_cell(root, f"I{base}", item.rule.unit_price)
            self._set_cell(root, f"K{base}", "中国")
            self._set_cell(root, f"M{base}", destination_country)
            self._set_cell(root, f"P{base}", "河北邢台")
            self._set_cell(root, f"S{base}", "照章征税")
            self._set_cell(root, f"T{base}", item.net_weight)
            self._set_cell(root, f"U{base}", item.gross_weight)
            self._set_cell(root, f"V{base}", result.box_count if index == 0 else None)
            self._set_cell(root, f"D{base + 1}", item.rule.declaration_elements)
            self._set_cell(root, f"I{base + 1}", f"=I{base}*G{base + 2}")
            self._set_cell(root, f"G{base + 2}", item.quantity)
            self._set_cell(root, f"H{base + 2}", "套")
            self._set_cell(root, f"I{base + 2}", "美元")

        workbook_xml = ET.fromstring(source.read("xl/workbook.xml"))
        calc = workbook_xml.find(f"{{{_MAIN_NS}}}calcPr")
        if calc is None:
            calc = ET.SubElement(workbook_xml, f"{{{_MAIN_NS}}}calcPr")
        calc.set("calcMode", "auto")
        calc.set("fullCalcOnLoad", "1")
        calc.set("forceFullCalc", "1")

        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == sheet_path:
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif info.filename == "xl/workbook.xml":
                    payload = ET.tostring(
                        workbook_xml, encoding="utf-8", xml_declaration=True
                    )
                target.writestr(info, payload)
        source.close()
        return output.getvalue()

    @classmethod
    def _normalize_item_styles(cls, root: ET.Element, capacity: int) -> None:
        """Extend the template's correctly formatted second item block to all later slots."""

        for item_index in range(1, capacity):
            destination_base = 20 + item_index * 3
            for row_offset in range(3):
                source_row = 23 + row_offset
                destination_row = destination_base + row_offset
                for column_index in range(1, 23):
                    column = get_column_letter(column_index)
                    cls._copy_cell_style(
                        root,
                        f"{column}{source_row}",
                        f"{column}{destination_row}",
                    )

    @staticmethod
    def _copy_cell_style(root: ET.Element, source_ref: str, destination_ref: str) -> None:
        source = root.find(f".//{{{_MAIN_NS}}}c[@r='{source_ref}']")
        destination = root.find(f".//{{{_MAIN_NS}}}c[@r='{destination_ref}']")
        if source is None or destination is None:
            raise ValueError(
                f"内置报关模板缺少样式单元格 {source_ref} 或 {destination_ref}。"
            )
        style_id = source.get("s")
        if style_id is None:
            destination.attrib.pop("s", None)
        else:
            destination.set("s", style_id)

    @staticmethod
    def _set_cell(root: ET.Element, reference: str, value: Any) -> None:
        cell = root.find(f".//{{{_MAIN_NS}}}c[@r='{reference}']")
        if cell is None:
            raise ValueError(f"内置报关模板缺少单元格 {reference}。")
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)
        if value is None:
            return
        if isinstance(value, str) and value.startswith("="):
            formula = ET.SubElement(cell, f"{{{_MAIN_NS}}}f")
            formula.text = value[1:]
            return
        if isinstance(value, str) and not value.isdigit():
            cell.set("t", "inlineStr")
            inline = ET.SubElement(cell, f"{{{_MAIN_NS}}}is")
            text = ET.SubElement(inline, f"{{{_MAIN_NS}}}t")
            text.text = value
            return
        cell.set("t", "n")
        number = ET.SubElement(cell, f"{{{_MAIN_NS}}}v")
        number.text = CustomsDeclarationService._number_text(value)

    @staticmethod
    def _number_text(value: Any) -> str:
        if isinstance(value, Decimal):
            return format(value, "f")
        return str(value)

    @staticmethod
    def _report_sheet_path(archive: zipfile.ZipFile) -> str:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        report_id: str | None = None
        for sheet in workbook.findall(f".//{{{_MAIN_NS}}}sheet"):
            if sheet.get("name") == "报关单":
                report_id = sheet.get(f"{{{_REL_NS}}}id")
                break
        if not report_id:
            raise ValueError("内置报关模板缺少“报关单”工作表。")
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for relationship in relationships.findall(f"{{{_PACKAGE_REL_NS}}}Relationship"):
            if relationship.get("Id") == report_id:
                target = relationship.get("Target", "")
                return "xl/" + target.lstrip("/").removeprefix("xl/")
        raise ValueError("内置报关模板的工作表关系损坏。")

    def _read_template_contract_number(self, template_path: Path) -> str:
        workbook = load_workbook(
            template_path,
            read_only=True,
            data_only=True,
            keep_links=True,
        )
        try:
            value = workbook["报关单"]["A10"].value
        finally:
            workbook.close()
        contract = self._text(value)
        if not contract:
            raise ValueError("内置报关模板缺少合同号。")
        return contract

    @classmethod
    def _load_catalog(cls, path: Path) -> list[ProductRule]:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if payload.get("version") != 1:
            raise ValueError("不支持的商品目录版本。")
        rules: list[ProductRule] = []
        for item in payload.get("products", []):
            rules.append(
                ProductRule(
                    rule_id=str(item["id"]),
                    skus=frozenset(str(value).upper() for value in item.get("skus", [])),
                    standard_name=str(item["standard_name"]),
                    declaration_name=str(item["declaration_name"]),
                    model=str(item["model"]),
                    material=str(item["material"]),
                    usage=str(item["usage"]),
                    unit_weight=Decimal(str(item["unit_weight"] or "0")),
                    unit_price=Decimal(str(item["unit_price_35"] or "0")),
                    number_key=cls._number_key(str(item["standard_name"])),
                )
            )
        return rules

    @staticmethod
    def _open_input_workbook(
        file_name: str,
        content: bytes,
        errors: list[CustomsIssue],
        label: str,
    ):
        if not file_name.lower().endswith(".xlsx"):
            errors.append(
                CustomsIssue(
                    code="UNSUPPORTED_FILE_TYPE",
                    message=f"{label}只支持 .xlsx 文件。",
                    file=file_name,
                )
            )
            return None
        if len(content) > MAX_FILE_SIZE:
            errors.append(
                CustomsIssue(
                    code="FILE_TOO_LARGE",
                    message=f"{label}超过 20MB 限制。",
                    file=file_name,
                )
            )
            return None
        try:
            with zipfile.ZipFile(io.BytesIO(content), "r") as archive:
                entries = archive.infolist()
                if len(entries) > MAX_ZIP_ENTRIES:
                    raise ValueError("压缩包文件项过多")
                if sum(item.file_size for item in entries) > MAX_EXPANDED_SIZE:
                    raise ValueError("压缩包展开后超过 100MB")
                if any(
                    item.filename.startswith("/") or ".." in Path(item.filename).parts
                    for item in entries
                ):
                    raise ValueError("压缩包包含非法路径")
            return load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        except Exception:
            errors.append(
                CustomsIssue(
                    code="INVALID_WORKBOOK",
                    message=f"{label}不是有效、未加密的 .xlsx 工作簿。",
                    file=file_name,
                )
            )
            return None

    @staticmethod
    def _positive_decimal(
        value: Any,
        errors: list[CustomsIssue],
        file_name: str,
        sheet: str,
        row: int,
        field: str,
    ) -> Decimal | None:
        try:
            result = Decimal(str(value).strip())
            if not result.is_finite() or result <= 0:
                raise InvalidOperation
            return result
        except (InvalidOperation, ValueError, TypeError):
            errors.append(
                CustomsIssue(
                    code="INVALID_NUMBER",
                    message=f"{field}必须是大于 0 的数字。",
                    file=file_name,
                    sheet=sheet,
                    row=row,
                    field=field,
                )
            )
            return None

    @classmethod
    def _positive_integer(
        cls,
        value: Any,
        errors: list[CustomsIssue],
        file_name: str,
        sheet: str,
        row: int,
        field: str,
    ) -> int | None:
        decimal_value = cls._positive_decimal(
            value, errors, file_name, sheet, row, field
        )
        if decimal_value is None:
            return None
        integral = decimal_value.to_integral_value()
        if decimal_value != integral:
            errors.append(
                CustomsIssue(
                    code="INVALID_INTEGER",
                    message=f"{field}必须是整数。",
                    file=file_name,
                    sheet=sheet,
                    row=row,
                    field=field,
                )
            )
            return None
        return int(integral)

    @staticmethod
    def _shipping_speed(value: str) -> ShippingSpeed | None:
        normalized = re.sub(r"\s+", "", value).replace("（", "(").replace("）", ")")
        for speed, channels in _SHIPPING_CHANNELS.items():
            if any(
                normalized
                == re.sub(r"\s+", "", channel).replace("（", "(").replace("）", ")")
                for channel in channels
            ):
                return speed
        return None

    @staticmethod
    def _trade_mode(value: Any) -> TradeMode | None:
        text = CustomsDeclarationService._text(value)
        if re.fullmatch(r"110(?:\.0+)?", text):
            text = "0110"
        elif re.fullmatch(r"9810(?:\.0+)?", text):
            text = "9810"
        try:
            return TradeMode(text)
        except ValueError:
            return None

    @staticmethod
    def _header_key(value: Any) -> str:
        return re.sub(r"\s+", "", CustomsDeclarationService._text(value))

    @staticmethod
    def _find_header_index(header: dict[str, int], needle: str) -> int | None:
        if needle in header:
            return header[needle]
        return next((index for name, index in header.items() if needle in name), None)

    @classmethod
    def _required_header_indices(
        cls, header: tuple[Any, ...], required: tuple[str, ...]
    ) -> dict[str, int] | None:
        normalized = {cls._header_key(value): index for index, value in enumerate(header)}
        indices = {
            name: cls._find_header_index(normalized, name)
            for name in required
        }
        if any(index is None for index in indices.values()):
            return None
        return {name: int(index) for name, index in indices.items()}

    @staticmethod
    def _value_at(row: tuple[Any, ...], index: int) -> Any:
        return row[index] if index < len(row) else None

    @staticmethod
    def _text(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _normalize_name(value: str) -> str:
        text = value.lower().replace("×", "x").replace("*", "x")
        text = text.replace("米", "m").replace("厘米", "cm").replace("毫米", "mm")
        text = re.sub(r"(透明|乳白|米白|白色|黑色|棕色|灰色|白|黑|棕|灰)", "", text)
        return re.sub(r"[\s_—–]+", "", text)

    @staticmethod
    def _detect_product_type(cleaned: str) -> str | None:
        for product_type, aliases in _TYPE_ALIASES:
            if any(alias in cleaned for alias in aliases):
                return product_type
        return None

    @staticmethod
    def _number_key(value: str) -> tuple[str, ...]:
        normalized = CustomsDeclarationService._normalize_name(value)
        result: list[str] = []
        for number in re.findall(r"\d+(?:\.\d+)?", normalized):
            rendered = format(Decimal(number), "f")
            if "." in rendered:
                rendered = rendered.rstrip("0").rstrip(".")
            result.append(rendered)
        return tuple(result)

    @staticmethod
    def _category_keys() -> Iterable[tuple[ShippingSpeed, TradeMode]]:
        for speed in ShippingSpeed:
            for trade in (TradeMode.GENERAL, TradeMode.CROSS_BORDER):
                yield speed, trade

    @staticmethod
    def _row_limit_issue(file_name: str, sheet: str) -> CustomsIssue:
        return CustomsIssue(
            code="ROW_LIMIT_EXCEEDED",
            message=f"工作表超过 {MAX_SHEET_ROWS} 行限制。",
            file=file_name,
            sheet=sheet,
        )

    @staticmethod
    def _shanghai_today() -> date:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Asia/Shanghai")).date()
