from __future__ import annotations

import io
import zipfile
from collections import OrderedDict
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Iterable

from openpyxl import load_workbook

from amazon_ops.customs_declaration.models import (
    AggregatedItem,
    CategoryResult,
    CustomsIssue,
    ProductRule,
    ShippingSpeed,
    TradeMode,
)
from amazon_ops.customs_declaration.service import (
    CustomsDeclarationService,
    MAX_EXPANDED_SIZE,
    MAX_FILE_SIZE,
    MAX_SHEET_ROWS,
    MAX_ZIP_ENTRIES,
    TEMPLATE_CAPACITIES,
)

from .models import EuropeanCategoryPreview, EuropeanCustomsPreview


MONEY_QUANTUM = Decimal("0.01")
WEIGHT_QUANTUM = Decimal("0.01")
CHANNELS = ("华贸卡航", "华贸海运", "华贸空运", "一八海运")
PREFIXES = ("红福", "其它")
COUNTRIES = ("德国", "法国", "英国")
TRANSPORTATION_MODES = {
    "华贸卡航": "公路运输",
    "华贸海运": "水路运输",
    "华贸空运": "航空运输",
    "一八海运": "水路运输",
}


@dataclass(slots=True)
class _EuropeanItem:
    product_name: str
    rule: ProductRule
    quantity: Decimal
    unit_price: Decimal
    unit_weight: Decimal
    net_weight: Decimal
    gross_weight: Decimal = Decimal("0")

    @property
    def amount(self) -> Decimal:
        return (self.quantity * self.unit_price).quantize(
            MONEY_QUANTUM, rounding=ROUND_HALF_UP
        )


@dataclass(slots=True)
class _EuropeanShipment:
    shipment_date: Any
    sender: str
    fba_number: str
    original_channel: str
    original_country: str
    channel: str
    prefix: str
    country: str
    box_count: int
    gross_weight: Decimal
    items: list[_EuropeanItem] = field(default_factory=list)

    @property
    def net_weight(self) -> Decimal:
        return sum((item.net_weight for item in self.items), Decimal("0"))

    @property
    def amount(self) -> Decimal:
        return sum((item.amount for item in self.items), Decimal("0"))


@dataclass(slots=True)
class _Analysis:
    preview: EuropeanCustomsPreview
    categories: dict[tuple[str, str, str], list[_EuropeanShipment]]
    report_results: dict[tuple[str, str, str], CategoryResult]


class EuropeanCustomsGenerationError(ValueError):
    def __init__(self, preview: EuropeanCustomsPreview):
        super().__init__("欧洲报关数据校验未通过。")
        self.preview = preview


class EuropeanCustomsDeclarationService:
    """Deterministic European declaration generator using only the uploaded Sheet1."""

    def __init__(self) -> None:
        # 欧洲货件仅提供分类和商品数据；成品沿用既有 0110 / 9810 报关单模板。
        self.report_service = CustomsDeclarationService()

    def preview(
        self, *, file_name: str, content: bytes, declaration_date: date | None = None
    ) -> EuropeanCustomsPreview:
        return self._analyze(
            file_name=file_name,
            content=content,
            declaration_date=declaration_date or self._shanghai_today(),
        ).preview

    def generate(
        self, *, file_name: str, content: bytes, declaration_date: date | None = None
    ) -> tuple[str, bytes]:
        selected_date = declaration_date or self._shanghai_today()
        analysis = self._analyze(
            file_name=file_name, content=content, declaration_date=selected_date
        )
        if not analysis.preview.can_generate:
            raise EuropeanCustomsGenerationError(analysis.preview)

        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for category in self._category_keys():
                shipments = analysis.categories.get(category, [])
                if not shipments:
                    continue
                channel, prefix, country = category
                archive.writestr(
                    f"欧洲报关表-{channel}-{prefix}-{country}-{selected_date:%Y%m%d}.xlsx",
                    self.report_service._render_workbook(
                        analysis.report_results[category],
                        selected_date,
                        destination_country=country,
                        trade_country=country if prefix == "红福" else None,
                        arrival_country=country,
                        destination_port=country,
                        transportation_mode=TRANSPORTATION_MODES[channel],
                    ),
                )
        return f"{selected_date:%Y%m%d}-欧洲报关表.zip", output.getvalue()

    def _analyze(self, *, file_name: str, content: bytes, declaration_date: date) -> _Analysis:
        errors: list[CustomsIssue] = []
        warnings: list[CustomsIssue] = []
        categories = self._parse_sheet1(file_name, content, errors, warnings)
        report_results = self._build_report_results(categories, errors, file_name)
        previews = [
            self._to_preview(
                category,
                categories.get(category, []),
                report_results.get(category),
            )
            for category in self._category_keys()
        ]
        return _Analysis(
            preview=EuropeanCustomsPreview(
                declaration_date=declaration_date,
                can_generate=not errors and any(item.generates_file for item in previews),
                categories=previews,
                warnings=warnings,
                errors=errors,
            ),
            categories=categories,
            report_results=report_results,
        )

    def _parse_sheet1(
        self,
        file_name: str,
        content: bytes,
        errors: list[CustomsIssue],
        warnings: list[CustomsIssue],
    ) -> dict[tuple[str, str, str], list[_EuropeanShipment]]:
        workbook = self._open_input_workbook(file_name, content, errors)
        if workbook is None:
            return {}
        if "Sheet1" not in workbook.sheetnames:
            errors.append(self._issue("MISSING_SHEET", "上传文件缺少“Sheet1”工作表。", file_name, "Sheet1"))
            return {}
        sheet = workbook["Sheet1"]
        if sheet.max_row > MAX_SHEET_ROWS:
            errors.append(self._issue("ROW_LIMIT_EXCEEDED", f"工作表超过 {MAX_SHEET_ROWS} 行限制。", file_name, sheet.title))
            return {}
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        indices = self._required_headers(
            header,
            ("发货日期", "发货人", "FBA号", "渠道", "国家", "箱数", "总重KG", "产品名称", "数量", "报关单价", "重量"),
        )
        if indices is None:
            errors.append(self._issue("MISSING_HEADER", "Sheet1 缺少欧洲报关所需表头。", file_name, sheet.title))
            return {}

        categories: dict[tuple[str, str, str], list[_EuropeanShipment]] = {}
        current: _EuropeanShipment | None = None
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            fba_number = self._text(self._value_at(row, indices["FBA号"]))
            if fba_number:
                channel_raw = self._text(self._value_at(row, indices["渠道"]))
                country_raw = self._text(self._value_at(row, indices["国家"]))
                channel = self._channel(channel_raw)
                prefix, country = self._destination(country_raw)
                boxes = self._positive_integer(
                    self._value_at(row, indices["箱数"]), errors, file_name, sheet.title, row_number, "箱数"
                )
                gross = self._positive_decimal(
                    self._value_at(row, indices["总重KG"]), errors, file_name, sheet.title, row_number, "总重KG"
                )
                if channel is None:
                    errors.append(
                        self._issue(
                            "INVALID_CHANNEL",
                            "渠道必须包含华贸卡航、华贸海运、华贸空运或一八海运。",
                            file_name,
                            sheet.title,
                            row_number,
                            "渠道",
                        )
                    )
                if country is None:
                    errors.append(self._issue("INVALID_COUNTRY", "国家必须以德国、法国或英国结尾。", file_name, sheet.title, row_number, "国家"))
                if channel is None or country is None or boxes is None or gross is None:
                    current = None
                    continue
                current = _EuropeanShipment(
                    shipment_date=self._value_at(row, indices["发货日期"]),
                    sender=self._text(self._value_at(row, indices["发货人"])),
                    fba_number=fba_number,
                    original_channel=channel_raw,
                    original_country=country_raw,
                    channel=channel,
                    prefix=prefix,
                    country=country,
                    box_count=boxes,
                    gross_weight=gross,
                )
                categories.setdefault((channel, prefix, country), []).append(current)

            product_name = self._text(self._value_at(row, indices["产品名称"]))
            if not product_name:
                continue
            if current is None:
                errors.append(self._issue("MISSING_SHIPMENT_CONTEXT", "商品行前缺少有效的 FBA 货件信息。", file_name, sheet.title, row_number))
                continue
            quantity = self._positive_decimal(self._value_at(row, indices["数量"]), errors, file_name, sheet.title, row_number, "数量")
            unit_price = self._positive_decimal(self._value_at(row, indices["报关单价"]), errors, file_name, sheet.title, row_number, "报关单价")
            unit_weight = self._positive_decimal(self._value_at(row, indices["重量"]), errors, file_name, sheet.title, row_number, "重量")
            if quantity is None or unit_price is None or unit_weight is None:
                continue
            rule = self.report_service._match_product(product_name)
            if rule is None:
                errors.append(
                    self._issue(
                        "UNKNOWN_PRODUCT",
                        f"产品名称“{product_name}”未匹配到内置报关商品目录。",
                        file_name,
                        sheet.title,
                        row_number,
                        "产品名称",
                    )
                )
                continue
            current.items.append(
                _EuropeanItem(
                    product_name=product_name,
                    rule=rule,
                    quantity=quantity,
                    unit_price=unit_price,
                    unit_weight=unit_weight,
                    net_weight=(quantity * unit_weight).quantize(WEIGHT_QUANTUM, rounding=ROUND_HALF_UP),
                )
            )

        for category, shipments in categories.items():
            for shipment in shipments:
                if not shipment.items:
                    errors.append(self._issue("MISSING_PRODUCTS", f"FBA号 {shipment.fba_number} 没有商品明细。", file_name, "Sheet1", field="产品名称"))
                    continue
        return categories

    def _build_report_results(
        self,
        categories: dict[tuple[str, str, str], list[_EuropeanShipment]],
        errors: list[CustomsIssue],
        file_name: str,
    ) -> dict[tuple[str, str, str], CategoryResult]:
        results: dict[tuple[str, str, str], CategoryResult] = {}
        for category, shipments in categories.items():
            if not shipments or not all(shipment.items for shipment in shipments):
                continue
            channel, prefix, country = category
            grouped: OrderedDict[tuple[str, Decimal], dict[str, Any]] = OrderedDict()
            for shipment in shipments:
                for item in shipment.items:
                    key = (item.rule.rule_id, item.unit_price)
                    aggregate = grouped.setdefault(
                        key,
                        {
                            "rule": replace(item.rule, unit_price=item.unit_price),
                            "quantity": Decimal("0"),
                            "net_weight": Decimal("0"),
                            "amount": Decimal("0"),
                        },
                    )
                    aggregate["quantity"] += item.quantity
                    aggregate["net_weight"] += item.net_weight
                    aggregate["amount"] += item.amount

            trade_mode = TradeMode.CROSS_BORDER if prefix == "红福" else TradeMode.GENERAL
            if len(grouped) > TEMPLATE_CAPACITIES[trade_mode]:
                errors.append(
                    self._issue(
                        "TEMPLATE_CAPACITY_EXCEEDED",
                        f"{channel}-{prefix}-{country} 有 {len(grouped)} 个报关商品，"
                        f"超过 {trade_mode} 模板 {TEMPLATE_CAPACITIES[trade_mode]} 个商品容量。",
                        file_name,
                        "Sheet1",
                    )
                )
                continue

            total_gross = sum((shipment.gross_weight for shipment in shipments), Decimal("0")).quantize(
                WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
            )
            allocation_weight = sum((entry["net_weight"] for entry in grouped.values()), Decimal("0")).quantize(
                WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
            )
            if allocation_weight <= 0:
                errors.append(
                    self._issue(
                        "ZERO_NET_WEIGHT",
                        f"{channel}-{prefix}-{country} 的商品净重必须大于 0。",
                        file_name,
                        "Sheet1",
                    )
                )
                continue

            allocated_gross = Decimal("0")
            total_net = (total_gross * Decimal("0.92")).quantize(
                WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
            )
            allocated_net = Decimal("0")
            items: list[AggregatedItem] = []
            aggregates = list(grouped.values())
            for index, aggregate in enumerate(aggregates):
                if index == len(aggregates) - 1:
                    gross_weight = (total_gross - allocated_gross).quantize(
                        WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
                    )
                    net_weight = (total_net - allocated_net).quantize(
                        WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
                    )
                else:
                    gross_weight = (total_gross * aggregate["net_weight"] / allocation_weight).quantize(
                        WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
                    )
                    net_weight = (total_net * aggregate["net_weight"] / allocation_weight).quantize(
                        WEIGHT_QUANTUM, rounding=ROUND_HALF_UP
                    )
                    allocated_gross += gross_weight
                    allocated_net += net_weight
                items.append(
                    AggregatedItem(
                        rule=aggregate["rule"],
                        quantity=aggregate["quantity"],
                        net_weight=net_weight,
                        gross_weight=gross_weight,
                        amount=aggregate["amount"].quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP),
                    )
                )
            results[category] = CategoryResult(
                shipping_speed=ShippingSpeed.STANDARD,
                trade_mode=trade_mode,
                order_ids=[shipment.fba_number for shipment in shipments],
                shipment_count=len(shipments),
                box_count=sum(shipment.box_count for shipment in shipments),
                gross_weight=total_gross,
                net_weight=total_net,
                amount=sum((item.amount for item in items), Decimal("0")).quantize(
                    MONEY_QUANTUM, rounding=ROUND_HALF_UP
                ),
                items=items,
            )
        return results

    def _to_preview(
        self,
        category: tuple[str, str, str],
        shipments: list[_EuropeanShipment],
        report_result: CategoryResult | None,
    ) -> EuropeanCategoryPreview:
        channel, prefix, country = category
        return EuropeanCategoryPreview(
            channel=channel,
            prefix=prefix,
            country=country,
            trade_mode="9810" if prefix == "红福" else "0110",
            shipment_count=len(shipments),
            item_count=len(report_result.items) if report_result else sum(len(shipment.items) for shipment in shipments),
            box_count=report_result.box_count if report_result else sum(shipment.box_count for shipment in shipments),
            net_weight=report_result.net_weight if report_result else Decimal("0"),
            gross_weight=report_result.gross_weight if report_result else sum((shipment.gross_weight for shipment in shipments), Decimal("0")).quantize(WEIGHT_QUANTUM, rounding=ROUND_HALF_UP),
            amount=report_result.amount if report_result else sum((shipment.amount for shipment in shipments), Decimal("0")).quantize(MONEY_QUANTUM, rounding=ROUND_HALF_UP),
            generates_file=bool(report_result and report_result.items),
        )

    @staticmethod
    def _open_input_workbook(file_name: str, content: bytes, errors: list[CustomsIssue]):
        if not file_name.lower().endswith(".xlsx"):
            errors.append(EuropeanCustomsDeclarationService._issue("UNSUPPORTED_FILE_TYPE", "欧洲报关表只支持 .xlsx 文件。", file_name))
            return None
        if len(content) > MAX_FILE_SIZE:
            errors.append(EuropeanCustomsDeclarationService._issue("FILE_TOO_LARGE", "文件超过 20MB 限制。", file_name))
            return None
        try:
            with zipfile.ZipFile(io.BytesIO(content), "r") as archive:
                entries = archive.infolist()
                if len(entries) > MAX_ZIP_ENTRIES or sum(item.file_size for item in entries) > MAX_EXPANDED_SIZE:
                    raise ValueError
            return load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        except Exception:
            errors.append(EuropeanCustomsDeclarationService._issue("INVALID_WORKBOOK", "文件不是有效、未加密的 .xlsx 工作簿。", file_name))
            return None

    @staticmethod
    def _required_headers(header: tuple[Any, ...], required: tuple[str, ...]) -> dict[str, int] | None:
        normalized = {EuropeanCustomsDeclarationService._header_key(value): index for index, value in enumerate(header)}
        indices: dict[str, int] = {}
        for name in required:
            index = normalized.get(name)
            if index is None:
                index = next((value for key, value in normalized.items() if name in key), None)
            if index is None:
                return None
            indices[name] = index
        return indices

    @staticmethod
    def _positive_decimal(value: Any, errors: list[CustomsIssue], file_name: str, sheet: str, row: int, field: str) -> Decimal | None:
        try:
            result = Decimal(str(value).strip())
            if not result.is_finite() or result <= 0:
                raise InvalidOperation
            return result
        except (InvalidOperation, ValueError, TypeError):
            errors.append(EuropeanCustomsDeclarationService._issue("INVALID_NUMBER", f"{field}必须是大于 0 的数字。", file_name, sheet, row, field))
            return None

    @classmethod
    def _positive_integer(cls, value: Any, errors: list[CustomsIssue], file_name: str, sheet: str, row: int, field: str) -> int | None:
        number = cls._positive_decimal(value, errors, file_name, sheet, row, field)
        if number is None:
            return None
        if number != number.to_integral_value():
            errors.append(cls._issue("INVALID_INTEGER", f"{field}必须是整数。", file_name, sheet, row, field))
            return None
        return int(number)

    @staticmethod
    def _channel(value: str) -> str | None:
        return next((channel for channel in CHANNELS if channel in value), None)

    @staticmethod
    def _destination(value: str) -> tuple[str, str | None]:
        country = next((item for item in COUNTRIES if value.endswith(item)), None)
        return ("红福" if value.startswith("红福") else "其它", country)

    @staticmethod
    def _category_keys() -> Iterable[tuple[str, str, str]]:
        for channel in CHANNELS:
            for prefix in PREFIXES:
                for country in COUNTRIES:
                    yield channel, prefix, country

    @staticmethod
    def _value_at(row: tuple[Any, ...], index: int) -> Any:
        return row[index] if index < len(row) else None

    @staticmethod
    def _header_key(value: Any) -> str:
        return "" if value is None else "".join(str(value).split())

    @staticmethod
    def _text(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _issue(code: str, message: str, file_name: str, sheet: str | None = None, row: int | None = None, field: str | None = None) -> CustomsIssue:
        return CustomsIssue(code=code, message=message, file=file_name, sheet=sheet, row=row, field=field)

    @staticmethod
    def _shanghai_today() -> date:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Asia/Shanghai")).date()
