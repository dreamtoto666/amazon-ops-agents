from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from amazon_ops.customs_declaration.models import ShippingSpeed, TradeMode
from amazon_ops.customs_declaration.service import (
    CustomsDeclarationService,
    CustomsGenerationError,
)


def workbook_bytes(sheet_name: str, rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def shipment_workbook(rows: list[list[object]]) -> bytes:
    return workbook_bytes(
        "美国专线箱单",
        [
            [
                "日期",
                "运营",
                "时效",
                "客户原单号\n（必填）",
                "FBA号",
                "仓库",
                "送货时间",
                "客户代号",
                "走货渠道\n（必填）",
                "报关方式",
                "贸易方式\n（单独报关必填）",
            ],
            *rows,
        ],
    )


def fba_workbook(
    rows: list[list[object]], *, detail_sheet_name: str = "货件详情"
) -> bytes:
    """Build an FBA workbook where packing product values are deliberately unusable."""

    workbook = Workbook()
    detail = workbook.active
    detail.title = detail_sheet_name
    detail_headers = [None] * 32
    detail_headers[0] = "货件单号"
    detail_headers[11] = "总申报量"
    detail_headers[17] = "品名"
    detail_headers[18] = "SKU"
    detail_headers[19] = "申报量"
    detail.append(detail_headers)
    packing = workbook.create_sheet("装箱明细")
    packing.append(["货件单号", "总箱数", "总重量", "品名", "SKU", "申报量", "箱号(箱库存)"])
    for row in rows:
        packing.append([
            row[0], row[1], row[2], "装箱表错误品名", "WRONG-SKU", 9999, row[6]
        ])
        detail_row = [None] * 32
        detail_row[0] = row[0]
        detail_row[17] = row[3]
        detail_row[18] = row[4]
        detail_row[19] = row[5]
        detail.append(detail_row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture(scope="module")
def service() -> CustomsDeclarationService:
    return CustomsDeclarationService()


def analyze(
    service: CustomsDeclarationService,
    shipment: bytes,
    fba: bytes,
):
    return service.preview(
        shipment_file_name="一八发货模板.xlsx",
        shipment_content=shipment,
        fba_file_name="FBA货件.xlsx",
        fba_content=fba,
        declaration_date=date(2026, 9, 1),
    )


@pytest.mark.parametrize(
    ("channel", "expected"),
    [
        ("美西普船快递派统配", ShippingSpeed.STANDARD),
        ("美西普船卡派统配（包）", ShippingSpeed.STANDARD),
        ("美东普船限时达纽约卡派（包）", ShippingSpeed.STANDARD),
        ("美东快船海铁纽约卡派（包）", ShippingSpeed.EXPRESS),
        ("美西快船卡派（包）", ShippingSpeed.EXPRESS),
        ("美西快船快递派", ShippingSpeed.EXPRESS),
        ("美森极致达卡派（包）", ShippingSpeed.MATSON),
        ("美森极致达快递派", ShippingSpeed.MATSON),
        ("美东普船限时达卡派特惠（包）", ShippingSpeed.STANDARD_ECONOMY),
        ("美西普船卡派特惠（包）", ShippingSpeed.STANDARD_ECONOMY),
    ],
)
def test_shipping_channel_groups_are_classified_exactly(
    channel: str,
    expected: ShippingSpeed,
) -> None:
    assert CustomsDeclarationService._shipping_speed(channel) == expected


def test_preview_has_eight_channel_trade_categories(
    service: CustomsDeclarationService,
) -> None:
    preview = analyze(
        service,
        shipment_workbook(
            [[None, None, None, "FBA-1", None, None, None, None, "美森极致达快递派", None, "9810"]]
        ),
        fba_workbook([["FBA-1", 1, 2, "D型-3米白", "F00398", 1, None]]),
    )

    assert len(preview.categories) == 8
    assert next(
        category
        for category in preview.categories
        if category.shipping_speed == ShippingSpeed.MATSON
        and category.trade_mode == TradeMode.CROSS_BORDER
    ).generates_file is True


@pytest.mark.parametrize(
    ("product_name", "standard_name", "unit_price"),
    [
        ("发泡D-10x10mm-3.05M", "发泡形车用密封条10x10mm-3m", Decimal("3.85")),
        ("发泡D-10x10mm-6.1M", "发泡形车用密封条10x10mm-6m", Decimal("5.6")),
        ("发泡D-15x15mm-3.05M", "发泡形车用密封条15x15mm-3m", Decimal("5.25")),
        ("发泡D-15x15mm-6.1M", "发泡形车用密封条15x15mm-6m", Decimal("7.35")),
        ("发泡D-19x19mm-3.05M", "发泡形车用密封条19x19mm-3m", Decimal("8.05")),
        ("发泡D-19x19mm-6.1M", "发泡形车用密封条19x19mm-6m", Decimal("12.25")),
        ("大五孔-15x12mm-5.03M", "五孔形车用密封条5.03m", Decimal("3.15")),
        ("大五孔-15x12mm-10.36M", "五孔形车用密封条10.36m", Decimal("5.25")),
        ("大五孔-15x12mm-16.4M", "五孔形车用密封条16.4m", Decimal("7.35")),
        ("大五孔-15x12mm-21.5M", "五孔形车用密封条21.5m", Decimal("9.1")),
        ("大五孔-15x12mm-26.5M", "五孔形车用密封条26.5m", Decimal("10.5")),
    ],
)
def test_confirmed_product_aliases_use_existing_rules_and_prices(
    service: CustomsDeclarationService,
    product_name: str,
    standard_name: str,
    unit_price: Decimal,
) -> None:
    rule = service._match_product(product_name)

    assert rule is not None
    assert rule.standard_name == standard_name
    assert rule.unit_price == unit_price


def test_d_column_classification_duplicate_rows_and_name_priority(
    service: CustomsDeclarationService,
) -> None:
    shipment = shipment_workbook(
        [
            [None, None, None, "FBA-1", "WRONG-E", None, None, None, "美西普船卡派统配（包）", None, "0110"],
            [None, None, None, "FBA-1", "WRONG-E", None, None, None, "美西普船卡派统配（包）", None, "0110"],
        ]
    )
    fba = fba_workbook(
        [
            ["FBA-1", 2, 10, "D型-7米黑", "F00499", 3, None],
            [None, None, None, "D型-7米白", "F00037", 2, None],
        ],
        detail_sheet_name="货件详细",
    )

    preview = analyze(service, shipment, fba)

    assert preview.can_generate is True
    category = preview.categories[0]
    assert (category.shipping_speed, category.trade_mode) == (
        ShippingSpeed.STANDARD,
        TradeMode.GENERAL,
    )
    assert category.order_count == 1
    assert category.shipment_count == 1
    assert category.item_count == 1
    assert category.items[0].quantity == Decimal("5")
    assert category.items[0].declaration_name == "D形门窗密封条"
    assert category.items[0].hs_code == "3926909090"
    assert any(issue.code == "SKU_NAME_CONFLICT" for issue in preview.warnings)


def test_product_type_size_mapping_epdm_hs_and_gross_tail(
    service: CustomsDeclarationService,
) -> None:
    shipment = shipment_workbook(
        [[None, None, None, "FBA-2", None, None, None, None, "美东快船海铁纽约卡派（包）", None, 9810]]
    )
    fba = fba_workbook(
        [
            ["FBA-2", 3, 20, "带导轨门底密封条2.44M", "F00499", 2, None],
            [None, None, None, "防霉条6*10M 白", "F00533", 3, None],
        ]
    )

    preview = analyze(service, shipment, fba)

    assert preview.can_generate is True
    category = preview.categories[3]
    assert category.box_count == 3
    assert category.gross_weight == Decimal("20.00")
    assert sum((item.gross_weight for item in category.items), Decimal("0")) == Decimal("20.00")
    epdm = next(item for item in category.items if item.material == "EPDM")
    pvc = next(item for item in category.items if item.material == "PVC")
    assert epdm.hs_code == "4016939000"
    assert pvc.hs_code == "3926909090"


def test_cross_category_order_is_blocking(service: CustomsDeclarationService) -> None:
    shipment = shipment_workbook(
        [
            [None, None, None, "FBA-X", None, None, None, None, "美西普船快递派统配", None, "0110"],
            [None, None, None, "FBA-X", None, None, None, None, "美西快船快递派", None, "9810"],
        ]
    )
    fba = fba_workbook([["FBA-X", 1, 1, "D型-3米白", "F00398", 1, None]])

    preview = analyze(service, shipment, fba)

    assert preview.can_generate is False
    assert any(issue.code == "ORDER_CATEGORY_CONFLICT" for issue in preview.errors)


def test_unknown_product_and_missing_shipment_are_blocking(
    service: CustomsDeclarationService,
) -> None:
    shipment = shipment_workbook(
        [
            [None, None, None, "FBA-FOUND", None, None, None, None, "美西普船快递派统配", None, "0110"],
            [None, None, None, "FBA-MISSING", None, None, None, None, "美西普船快递派统配", None, "0110"],
        ]
    )
    fba = fba_workbook([["FBA-FOUND", 1, 2, "不存在的商品-1米", "UNKNOWN", 1, None]])

    preview = analyze(service, shipment, fba)

    assert preview.can_generate is False
    assert {issue.code for issue in preview.errors} >= {
        "UNKNOWN_PRODUCT",
        "MISSING_SHIPMENT",
    }


@pytest.mark.parametrize(
    ("trade_mode", "capacity_exceeding_count"),
    [("0110", 24), ("9810", 17)],
)
def test_template_capacity_is_blocking(
    service: CustomsDeclarationService,
    trade_mode: str,
    capacity_exceeding_count: int,
) -> None:
    shipment = shipment_workbook(
        [[None, None, None, "FBA-CAP", None, None, None, None, "美西普船快递派统配", None, trade_mode]]
    )
    selected_rules = [rule for rule in service.rules if rule.unit_weight > 0][:capacity_exceeding_count]
    fba_rows: list[list[object]] = []
    for index, rule in enumerate(selected_rules):
        fba_rows.append(
            [
                "FBA-CAP" if index == 0 else None,
                1 if index == 0 else None,
                100 if index == 0 else None,
                rule.standard_name,
                next(iter(rule.skus), ""),
                1,
                None,
            ]
        )

    preview = analyze(service, shipment, fba_workbook(fba_rows))

    assert preview.can_generate is False
    assert any(issue.code == "TEMPLATE_CAPACITY_EXCEEDED" for issue in preview.errors)


def test_golden_97_box_summary(service: CustomsDeclarationService) -> None:
    shipment = shipment_workbook(
        [[None, None, None, "FBA-GOLD", None, None, None, None, "美西普船卡派统配（包）", None, "0110"]]
    )
    products = [
        ("D型-7米白", "F00037", 160),
        ("车库门侧密封条10.98米", "F00494", 510),
        ("车库门侧密封条12.15米", "F00524", 434),
        ("车库门侧密封条7米", "F00526", 36),
        ("带导轨门底密封条2.44M", "F00499", 108),
        ("镂空挡水条白-1.7米", "F00504", 128),
        ("镂空挡水条透明-1米", "F00507", 240),
        ("镂空挡水条透明-3.5米", "F00510", 64),
        ("平面墙角保护器-白2cm-3.3m", "F00460", 124),
        ("平面墙角保护器-白2cm-6.1m", "F00459", 84),
        ("平面墙角保护器-白4cm-6.1m", "F00457", 32),
        ("平面墙角保护器-白6cm-6.1m", "F00501", 80),
        ("扇形双面装饰条-3米", "F00514", 8),
        ("扇形双面装饰条-6米", "F00516", 16),
    ]
    fba_rows: list[list[object]] = []
    for index, (name, sku, quantity) in enumerate(products):
        fba_rows.append(
            ["FBA-GOLD" if index == 0 else None, 97 if index == 0 else None, 1913.89 if index == 0 else None, name, sku, quantity, None]
        )

    preview = analyze(service, shipment, fba_workbook(fba_rows))

    assert preview.errors == []
    category = preview.categories[0]
    assert category.item_count == 14
    assert category.box_count == 97
    assert category.net_weight == Decimal("1760.78")
    assert category.gross_weight == Decimal("1913.89")
    assert category.amount == Decimal("18901.40")
    assert sum((item.gross_weight for item in category.items), Decimal("0")) == Decimal("1913.89")
    assert sum((item.net_weight for item in category.items), Decimal("0")) == Decimal("1760.78")


def test_generate_preserves_template_and_fills_report(
    service: CustomsDeclarationService,
) -> None:
    shipment = shipment_workbook(
        [[None, None, None, "FBA-1", None, None, None, None, "美西普船快递派统配", None, "0110"]]
    )
    fba = fba_workbook([["FBA-1", 2, 10, "D型-7米白", "F00037", 5, None]])

    zip_name, archive_bytes = service.generate(
        shipment_file_name="一八发货模板.xlsx",
        shipment_content=shipment,
        fba_file_name="FBA货件.xlsx",
        fba_content=fba,
        declaration_date=date(2026, 9, 1),
    )

    assert zip_name == "20260901-报关单.zip"
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        names = archive.namelist()
        assert names == ["2608290072-美国海运-普船统配-0110-2箱-报关单-25美元.xlsx"]
        generated = archive.read(names[0])
    with zipfile.ZipFile(io.BytesIO(generated)) as output_zip, zipfile.ZipFile(
        service.template_path
    ) as template_zip:
        assert len([name for name in output_zip.namelist() if name.startswith("xl/media/")]) == len(
            [name for name in template_zip.namelist() if name.startswith("xl/media/")]
        )
        assert len([name for name in output_zip.namelist() if "drawing" in name]) == len(
            [name for name in template_zip.namelist() if "drawing" in name]
        )
    workbook = load_workbook(io.BytesIO(generated), data_only=False, read_only=True)
    assert workbook.sheetnames == ["报关单", "发票", "箱单", "合同", "申报要素", "填制规范"]
    report = workbook["报关单"]
    assert report["K4"].value.date() == date(2026, 9, 1)
    assert report["A10"].value == "2608290072"
    assert report["K12"].value == "USD/3113/3"
    assert report["B20"].value == 3926909090
    assert report["D20"].value == "D形门窗密封条"
    assert report["G22"].value == 5
    assert report["I21"].value == "=I20*G22"
    assert report["T20"].value == 9.2
    assert report["U20"].value == 10
    assert report["V20"].value == 2
    workbook.close()


def test_9810_template_is_selected_for_9810_category(
    service: CustomsDeclarationService,
) -> None:
    shipment = shipment_workbook(
        [[None, None, None, "FBA-9810", None, None, None, None, "美西快船快递派", None, "9810"]]
    )

    _, archive_bytes = service.generate(
        shipment_file_name="一八发货模板.xlsx",
        shipment_content=shipment,
        fba_file_name="FBA货件.xlsx",
        fba_content=fba_workbook([["FBA-9810", 1, 2, "D型-3米白", "F00398", 2, None]]),
        declaration_date=date(2026, 9, 1),
    )

    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        name = archive.namelist()[0]
        generated = archive.read(name)
    assert name == "2608290074-美国海运-快船-9810-1箱-报关单-6美元.xlsx"
    workbook = load_workbook(io.BytesIO(generated), data_only=False, read_only=True)
    try:
        report = workbook["报关单"]
        assert report["E8"].value == 9810
        assert report["K20"].value == "中国"
        assert report["M20"].value == "美国"
        assert report["P20"].value == "河北邢台"
        assert report["S20"].value == "照章征税"
    finally:
        workbook.close()


def test_later_product_rows_reuse_correct_style_and_fixed_fields(
    service: CustomsDeclarationService,
) -> None:
    shipment = shipment_workbook(
        [[None, None, None, "FBA-STYLE", None, None, None, None, "美西普船快递派统配", None, "0110"]]
    )
    selected_rules = [rule for rule in service.rules if rule.unit_weight > 0][:8]
    rows: list[list[object]] = []
    for index, rule in enumerate(selected_rules):
        rows.append(
            [
                "FBA-STYLE" if index == 0 else None,
                2 if index == 0 else None,
                20 if index == 0 else None,
                rule.standard_name,
                next(iter(rule.skus), ""),
                1,
                None,
            ]
        )

    _, archive_bytes = service.generate(
        shipment_file_name="一八发货模板.xlsx",
        shipment_content=shipment,
        fba_file_name="FBA货件.xlsx",
        fba_content=fba_workbook(rows),
        declaration_date=date(2026, 9, 1),
    )

    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        generated = archive.read(archive.namelist()[0])
    workbook = load_workbook(io.BytesIO(generated), data_only=False)
    report = workbook["报关单"]
    for base in (35, 38, 41):
        assert report[f"B{base}"]._style == report["B23"]._style
        assert report[f"D{base}"]._style == report["D23"]._style
        assert report[f"D{base + 1}"]._style == report["D24"]._style
        assert report[f"G{base + 2}"]._style == report["G25"]._style
        assert report[f"K{base}"].value == "中国"
        assert report[f"M{base}"].value == "美国"
        assert report[f"P{base}"].value == "河北邢台"
        assert report[f"S{base}"].value == "照章征税"
    workbook.close()


def test_generation_rejects_blocking_preview(service: CustomsDeclarationService) -> None:
    with pytest.raises(CustomsGenerationError):
        service.generate(
            shipment_file_name="bad.xls",
            shipment_content=b"not xlsx",
            fba_file_name="bad.xls",
            fba_content=b"not xlsx",
            declaration_date=date(2026, 9, 1),
        )
