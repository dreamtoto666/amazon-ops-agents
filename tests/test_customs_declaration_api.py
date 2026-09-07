from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook

from amazon_ops.api import create_app
from amazon_ops.auth import AuthUser
from amazon_ops.idempotency import InMemoryIdempotencyRegistry
from amazon_ops.memory import InMemoryConversationStore


class StubRunManager:
    def __init__(self) -> None:
        self.memory = InMemoryConversationStore()
        self.roles = type("Roles", (), {"llm": None})()

    def health(self):
        return {"status": "ok"}


class StubAuthStore:
    def user_for_token(self, token):
        if token == "test-token":
            return AuthUser("test-user", None, "operator", True, "operator")
        return None

    def close(self) -> None:
        return None


def workbook_bytes(sheet_name: str, rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def valid_files() -> dict[str, tuple[str, bytes, str]]:
    shipment = workbook_bytes(
        "美国专线箱单",
        [
            ["日期", "运营", "时效", "客户原单号", "FBA号", "仓库", "时间", "客户", "走货渠道", "报关方式", "贸易方式"],
            [None, None, None, "FBA-API", None, None, None, None, "美西普船快递派统配", None, "0110"],
        ],
    )
    fba_workbook = Workbook()
    detail = fba_workbook.active
    detail.title = "货件详情"
    detail_headers = [None] * 32
    detail_headers[0] = "货件单号"
    detail_headers[11] = "总申报量"
    detail_headers[17] = "品名"
    detail_headers[18] = "SKU"
    detail_headers[19] = "申报量"
    detail.append(detail_headers)
    detail_row = [None] * 32
    detail_row[0] = "FBA-API"
    detail_row[17] = "D型-3米白"
    detail_row[18] = "F00398"
    detail_row[19] = 2
    detail.append(detail_row)
    packing = fba_workbook.create_sheet("装箱明细")
    packing.append(["货件单号", "总箱数", "总重量", "品名", "SKU", "申报量"])
    packing.append(["FBA-API", 1, 2, "错误品名", "WRONG-SKU", 9999])
    fba_output = io.BytesIO()
    fba_workbook.save(fba_output)
    fba = fba_output.getvalue()
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return {
        "shipment_file": ("一八发货模板.xlsx", shipment, media_type),
        "fba_file": ("FBA货件.xlsx", fba, media_type),
    }


def client(*, authenticated: bool = True) -> TestClient:
    headers = {"Authorization": "Bearer test-token"} if authenticated else {}
    return TestClient(
        create_app(
            StubRunManager(),
            idempotency_registry=InMemoryIdempotencyRegistry(),
            auth_store=StubAuthStore(),
        ),
        headers=headers,
    )


def test_customs_preview_requires_authentication() -> None:
    response = client(authenticated=False).post(
        "/api/customs-declarations/preview", files=valid_files()
    )

    assert response.status_code == 401


def test_customs_preview_and_zip_generation() -> None:
    api = client()

    preview = api.post("/api/customs-declarations/preview", files=valid_files())
    generated = api.post("/api/customs-declarations/generate", files=valid_files())

    assert preview.status_code == 200
    assert preview.json()["can_generate"] is True
    assert len(preview.json()["categories"]) == 8
    assert generated.status_code == 200
    assert generated.headers["content-type"] == "application/zip"
    assert generated.content.startswith(b"PK")


def test_customs_generate_returns_structured_422() -> None:
    files = valid_files()
    files["fba_file"] = ("bad.xls", b"invalid", "application/vnd.ms-excel")

    response = client().post("/api/customs-declarations/generate", files=files)

    assert response.status_code == 422
    assert response.json()["can_generate"] is False
    assert response.json()["errors"][0]["code"] == "UNSUPPORTED_FILE_TYPE"
