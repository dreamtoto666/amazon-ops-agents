from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from amazon_ops.european_customs_declaration.service import (
    EuropeanCustomsDeclarationService,
    EuropeanCustomsGenerationError,
)


def workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["发货日期", "发货人", "FBA号", "渠道", "国家", "箱数", "总重KG", "产品名称", "数量", "报关单价", "重量"])
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture(scope="module")
def service() -> EuropeanCustomsDeclarationService:
    return EuropeanCustomsDeclarationService()


def test_fixed_24_categories_normalization_and_weight_allocation(
    service: EuropeanCustomsDeclarationService,
) -> None:
    content = workbook_bytes([
        [date(2026, 9, 1), "甲", "FBA-A", "华贸海运", "红福德国", 2, 10, "挡水条1m", 2, 5, 1],
        [None, None, None, None, None, None, None, "挡水条1.5m", 3, 4, 2],
        [date(2026, 9, 1), "乙", "FBA-B", "华贸卡航", "其它法国", 1, 4, "挡水条1m", 1, 6, 1],
    ])

    preview = service.preview(file_name="欧洲货件.xlsx", content=content, declaration_date=date(2026, 9, 3))

    assert preview.can_generate is True
    assert len(preview.categories) == 24
    assert sum(category.generates_file for category in preview.categories) == 2
    sea_red_germany = next(category for category in preview.categories if (category.channel, category.prefix, category.country) == ("华贸海运", "红福", "德国"))
    assert sea_red_germany.trade_mode == "9810"
    assert sea_red_germany.item_count == 2
    assert sea_red_germany.net_weight == Decimal("9.20")
    assert sea_red_germany.gross_weight == Decimal("10.00")
    assert sea_red_germany.amount == Decimal("22.00")
    assert next(category for category in preview.categories if category.prefix == "其它").trade_mode == "0110"


def test_generate_existing_declaration_template_with_country_and_totals(
    service: EuropeanCustomsDeclarationService,
) -> None:
    content = workbook_bytes([
        [date(2026, 9, 1), "甲", "FBA-A", "华贸海运", "红福德国", 2, 10, "挡水条1m", 2, 5, 1],
        [None, None, None, None, None, None, None, "挡水条1.5m", 3, 4, 2],
    ])

    zip_name, payload = service.generate(file_name="欧洲货件.xlsx", content=content, declaration_date=date(2026, 9, 3))

    assert zip_name == "20260903-欧洲报关表.zip"
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.namelist() == ["欧洲报关表-华贸海运-红福-德国-20260903.xlsx"]
        generated = archive.read(archive.namelist()[0])
    workbook = load_workbook(io.BytesIO(generated), data_only=False, read_only=True)
    try:
        assert workbook.sheetnames == ["报关单", "发票", "箱单", "合同", "申报要素", "填制规范"]
        sheet = workbook["报关单"]
        assert sheet["E8"].value == 9810
        assert sheet["E6"].value == "水路运输"
        assert sheet["E10"].value == "德国"
        assert sheet["G10"].value == "德国"
        assert sheet["K10"].value == "德国"
        assert sheet["B20"].value == 3926909090
        assert sheet["D20"].value == "挡水条"
        assert sheet["I20"].value == 5
        assert sheet["K20"].value == "中国"
        assert sheet["M20"].value == "德国"
        assert sheet["P20"].value == "河北邢台"
        assert sheet["S20"].value == "照章征税"
        assert sheet["T20"].value == 2.3
        assert sheet["U20"].value == 2.5
        assert sheet["V20"].value == 2
        assert sheet["T23"].value == 6.9
        assert sheet["U23"].value == 7.5
        assert sheet["D21"].value.startswith("0|0|")
        assert sheet["G22"].value == 2
        assert sheet["I21"].value == "=I20*G22"
    finally:
        workbook.close()


def test_other_prefix_uses_0110_template_and_destination_country(
    service: EuropeanCustomsDeclarationService,
) -> None:
    content = workbook_bytes([
        [date(2026, 9, 1), "乙", "FBA-B", "华贸空运", "其它法国", 1, 4, "挡水条1m", 1, 6, 1],
    ])

    _, payload = service.generate(
        file_name="欧洲货件.xlsx", content=content, declaration_date=date(2026, 9, 3)
    )

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        generated = archive.read("欧洲报关表-华贸空运-其它-法国-20260903.xlsx")
    workbook = load_workbook(io.BytesIO(generated), data_only=False, read_only=True)
    try:
        sheet = workbook["报关单"]
        assert sheet["E8"].value == "0110"
        assert sheet["E6"].value == "航空运输"
        assert sheet["E10"].value == "香港"
        assert sheet["G10"].value == "法国"
        assert sheet["K10"].value == "法国"
        assert sheet["M20"].value == "法国"
        assert sheet["I20"].value == 6
    finally:
        workbook.close()


def test_truck_channel_uses_road_transport(
    service: EuropeanCustomsDeclarationService,
) -> None:
    content = workbook_bytes([
        [date(2026, 9, 1), "丙", "FBA-C", "华贸卡航", "红福英国", 1, 4, "挡水条1m", 1, 6, 1],
    ])

    _, payload = service.generate(
        file_name="欧洲货件.xlsx", content=content, declaration_date=date(2026, 9, 3)
    )

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        generated = archive.read("欧洲报关表-华贸卡航-红福-英国-20260903.xlsx")
    workbook = load_workbook(io.BytesIO(generated), data_only=False, read_only=True)
    try:
        assert workbook["报关单"]["E6"].value == "公路运输"
    finally:
        workbook.close()


def test_one_eight_sea_channel_is_a_separate_category(
    service: EuropeanCustomsDeclarationService,
) -> None:
    content = workbook_bytes([
        [date(2026, 9, 1), "丁", "FBA-D", "一八海运", "其它英国", 1, 4, "挡水条1m", 1, 6, 1],
    ])

    preview = service.preview(
        file_name="欧洲货件.xlsx", content=content, declaration_date=date(2026, 9, 3)
    )
    one_eight = next(
        category
        for category in preview.categories
        if (category.channel, category.prefix, category.country) == ("一八海运", "其它", "英国")
    )
    assert one_eight.generates_file is True

    _, payload = service.generate(
        file_name="欧洲货件.xlsx", content=content, declaration_date=date(2026, 9, 3)
    )
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        generated = archive.read("欧洲报关表-一八海运-其它-英国-20260903.xlsx")
    workbook = load_workbook(io.BytesIO(generated), data_only=False, read_only=True)
    try:
        assert workbook["报关单"]["E6"].value == "水路运输"
    finally:
        workbook.close()


def test_invalid_weight_blocks_generation(service: EuropeanCustomsDeclarationService) -> None:
    content = workbook_bytes([[date(2026, 9, 1), "甲", "FBA-A", "华贸海运", "红福德国", 1, 1, "挡水条1m", 1, 2, None]])

    preview = service.preview(file_name="欧洲货件.xlsx", content=content)
    assert preview.can_generate is False
    assert any(issue.field == "重量" for issue in preview.errors)
    with pytest.raises(EuropeanCustomsGenerationError):
        service.generate(file_name="欧洲货件.xlsx", content=content)
